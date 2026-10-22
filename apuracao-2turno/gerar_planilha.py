"""Gera apuracao_2turno.xlsx: painel de acompanhamento do 2º turno (TSE, arquivo EA20).

Uso: python3 gerar_planilha.py [saida.xlsx]

A planilha sai vazia, com Lula (13) e Flávio Bolsonaro (22) já cadastrados; os votos
chegam pela conexão com o TSE (Power Query no Excel ou Apps Script no Google Planilhas).
"""
import math
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = Path(__file__).parent
HIST_LINHAS = 1000  # linhas de histórico com fórmulas (1 por minuto ≈ 16 h)

AZUL = "1F4E79"
AZUL_CLARO = "DDEBF7"
COR1 = "C00000"  # candidato 1 (menor número de urna)
COR2 = "2E75B6"  # candidato 2
CINZA = "F2F2F2"

TITULO = Font(name="Calibri", size=18, bold=True, color="FFFFFF")
CAB = Font(bold=True, color="FFFFFF")
NEGRITO = Font(bold=True)
FUNDO_AZUL = PatternFill("solid", fgColor=AZUL)
FUNDO_CLARO = PatternFill("solid", fgColor=AZUL_CLARO)
FUNDO_CINZA = PatternFill("solid", fgColor=CINZA)
FINO = Side(style="thin", color="BFBFBF")
BORDA = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
QUEBRA = Alignment(wrap_text=True, vertical="top")

CAMPOS_RESUMO = [
    ("ele", "Código da eleição"), ("t", "Turno"), ("tpabr", "Tipo de abrangência"),
    ("cdabr", "Código da abrangência"), ("cargo", "Cargo"), ("dt", "Data da totalização"),
    ("ht", "Hora da totalização"), ("and", "Andamento (n/p/f)"),
    ("md", "Matematicamente definido (e/s/n)"), ("tf", "Totalização final (s/n)"),
    ("dv", "Divulga votação (s/n)"), ("ts", "Seções"), ("st", "Seções totalizadas"),
    ("te", "Eleitorado"), ("est", "Eleitorado em seções totalizadas"),
    ("esi", "Eleitorado em seções instaladas"), ("c", "Comparecimento"), ("a", "Abstenção"),
    ("tv", "Total de votos"), ("vv", "Votos válidos"), ("vb", "Votos brancos"),
    ("tvn", "Total de votos nulos"), ("van", "Votos anulados"), ("vansj", "Anulados sub judice"),
]
CAB_HIST = ["Data", "Hora", "Secoes", "SecoesTot", "Eleitorado", "EleitTot", "EleitInst",
            "Comparec", "VotosTotais", "VotosValidos", "VotosCand1", "VotosCand2"]


def largura(ws, larguras):
    for col, w in larguras.items():
        ws.column_dimensions[col].width = w


def faixa_titulo(ws, rng, texto):
    ws.merge_cells(rng)
    c = ws[rng.split(":")[0]]
    c.value = texto
    c.font = TITULO
    c.fill = FUNDO_AZUL
    c.alignment = Alignment(vertical="center", indent=1)


def aba_config(wb):
    ws = wb.active
    ws.title = "Config"
    largura(ws, {"A": 44, "B": 70, "C": 60})
    faixa_titulo(ws, "A1:C1", "Configuração da conexão com o TSE")
    ws.row_dimensions[1].height = 32

    linhas = [
        (3, "Endereço base (oficial)", "https://resultados.tse.jus.br/oficial",
         "Simulado do TSE: https://resultados-sim.tse.jus.br/simulado/simulado2026"),
        (4, "Pasta do pleito", "ele2026", ""),
        (5, "Código da eleição (2º turno 2026)", 6258,
         "1º turno = 6257 · 2º turno = 6258 (código usado no nome do arquivo com 6 dígitos)"),
        (6, "Abrangência (br ou sigla da UF)", "br",
         "br = Brasil (Presidente). Para Governador use a UF, ex.: sp"),
        (7, "Código do cargo", 1, "1 = Presidente · 3 = Governador"),
        (8, "URL manual (opcional — tem prioridade)", None,
         "Cole aqui a URL exata se o TSE publicar em outro caminho"),
    ]
    for r, rot, val, obs in linhas:
        ws.cell(r, 1, rot).font = NEGRITO
        c = ws.cell(r, 2, val)
        c.fill = FUNDO_CLARO
        c.border = BORDA
        ws.cell(r, 3, obs).alignment = QUEBRA

    dv = DataValidation(type="list", formula1='"1,3"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B7")

    ws["A10"] = "Formato do nome (spec. EA20)"
    ws["A10"].font = NEGRITO
    ws["B10"] = "<abrangência>-c<cargo 4 díg.>-e<eleição 6 díg.>-u.json   ex.: br-c0001-e006258-u.json"
    ws["A12"] = "URL usada na conexão"
    ws["A12"].font = Font(bold=True, color=AZUL)
    ws["B12"] = ('=IF(B8<>"",B8,B3&"/"&B4&"/"&B5&"/dados/"&LOWER(B6)&"/"&LOWER(B6)'
                 '&"-c"&TEXT(B7,"0000")&"-e"&TEXT(B5,"000000")&"-u.json")')
    ws["B12"].font = Font(bold=True, color=AZUL)
    ws["B12"].fill = FUNDO_CINZA
    ws["C12"] = "Nome definido URL_TSE (lido pelo Power Query / Apps Script)"
    wb.defined_names["URL_TSE"] = DefinedName("URL_TSE", attr_text="Config!$B$12")

    ws["A15"] = "Parâmetros do modelo de probabilidade"
    ws["A15"].font = Font(bold=True, size=13, color=AZUL)
    params = [
        (16, "Incerteza dos votos restantes (desvio-padrão)", 0.06, "SIGMA",
         "Quanto a fatia do candidato nas urnas que faltam pode diferir da fatia atual. "
         "Regiões apuram em ritmos diferentes (ex.: 2022, virada no fim). 4%–8% é razoável."),
        (17, "Ajuste de tendência p/ Lula (p.p.)", 0.0, "AJUSTE",
         "Ex.: +0,02 se você espera que as urnas restantes favoreçam Lula em 2 p.p. (negativo favorece Flávio)"),
        (18, "Comparecimento padrão (antes dos 1ºs dados)", 0.79, "COMPAREC_PAD",
         "Usado só enquanto não há seções totalizadas"),
        (19, "Fração de votos válidos padrão", 0.95, "VALIDOS_PAD",
         "Usado só enquanto não há votos totalizados"),
    ]
    for r, rot, val, nome, obs in params:
        ws.cell(r, 1, rot).font = NEGRITO
        c = ws.cell(r, 2, val)
        c.number_format = "0.0%"
        c.fill = FUNDO_CLARO
        c.border = BORDA
        ws.cell(r, 3, obs).alignment = QUEBRA
        ws.row_dimensions[r].height = 45
        wb.defined_names[nome] = DefinedName(nome, attr_text=f"Config!$B${r}")

    ws["A22"] = "Atenção"
    ws["A22"].font = Font(bold=True, color=COR1)
    ws["B22"] = ("Lula (13) e Flávio Bolsonaro (22) já estão cadastrados; os votos chegam pela conexão "
                 "(aba 'Como usar'). Para acompanhar sem configurar nada, use o painel web da pasta "
                 "painel-ao-vivo. A probabilidade é uma estimativa; o resultado oficial é sempre o do TSE.")
    ws["B22"].alignment = QUEBRA
    ws.row_dimensions[22].height = 48


def aba_dados(wb):
    ws = wb.create_sheet("Dados")
    largura(ws, {"A": 12, "B": 16, "C": 34, "D": 3, "E": 9, "F": 26, "G": 12, "H": 14,
                 "I": 8, "J": 12, "K": 6})
    for col, txt in zip("AB", ["Campo", "Valor"]):
        ws[f"{col}1"] = txt
    for i, txt in enumerate(["Numero", "Nome", "Partido", "Votos", "Eleito", "Situacao", "Seq"]):
        ws.cell(1, 5 + i, txt)
    for c in list(ws[1]):
        if c.value:
            c.font = CAB
            c.fill = FUNDO_AZUL
    ws["C1"] = "Descrição (spec. EA20)"
    ws["C1"].font = NEGRITO
    for i, (k, desc) in enumerate(CAMPOS_RESUMO, start=2):
        ws.cell(i, 1, k)
        ws.cell(i, 3, desc).font = Font(italic=True, color="7F7F7F")
    ws["E30"] = ("Esta aba é preenchida automaticamente (Power Query em A1 e E1, ou Apps Script). "
                 "Não digite aqui.")
    ws["E30"].font = Font(italic=True, color="7F7F7F")
    return ws


CANDIDATOS = [[13, "LULA", "PT", 0, "n", "", 1],
              [22, "FLÁVIO BOLSONARO", "PL", 0, "n", "", 2]]


def preencher_inicial(ws_dados):
    """Antes da 1ª leitura do TSE: só os candidatos do 2º turno, sem votos."""
    iniciais = {"ele": "6258", "t": "2", "tpabr": "br", "cdabr": "br", "cargo": "Presidente",
                "and": "n", "md": "n", "tf": "n", "dv": "n"}
    for i, (k, _) in enumerate(CAMPOS_RESUMO, start=2):
        ws_dados.cell(i, 2, iniciais.get(k, 0 if k not in ("dt", "ht") else ""))
    for r, cand in enumerate(CANDIDATOS, start=2):
        for j, v in enumerate(cand):
            ws_dados.cell(r, 5 + j, v)


def formula_prob(r):
    """Probabilidade do candidato 1 na linha r do Histórico (colunas calculadas N:U)."""
    return (f'=IF($A{r}="",NA(),IF(R{r}<=0,IF(K{r}>L{r},1,IF(K{r}<L{r},0,0.5)),'
            f'1-NORMDIST(S{r},P{r}+AJUSTE,SIGMA,TRUE)))')


def aba_historico(wb):
    ws = wb.create_sheet("Histórico")
    for j, txt in enumerate(CAB_HIST, start=1):
        c = ws.cell(1, j, txt)
        c.font = CAB
        c.fill = FUNDO_AZUL
    calc = ["% seções", "% eleitorado", "% válidos Lula", "% válidos Flávio",
            "Válidos restantes (estim.)", "Fatia que Lula precisa", "Prob. Lula", "Prob. Flávio"]
    for j, txt in enumerate(calc, start=14):
        c = ws.cell(1, j, txt)
        c.font = CAB
        c.fill = PatternFill("solid", fgColor="7F7F7F")
        c.alignment = CENTRO
    ws.row_dimensions[1].height = 32
    largura(ws, {"A": 11, "B": 9, **{ch: 12 for ch in "CDEFGHIJKL"}, "M": 3,
                 **{ch: 13 for ch in "NOPQRSTU"}})

    for r in range(2, HIST_LINHAS + 2):
        vazio = f'$A{r}=""'
        ws[f"N{r}"] = f"=IF({vazio},NA(),IF(C{r}>0,D{r}/C{r},0))"
        ws[f"O{r}"] = f"=IF({vazio},NA(),IF(E{r}>0,F{r}/E{r},0))"
        ws[f"P{r}"] = f"=IF({vazio},NA(),IF(K{r}+L{r}>0,K{r}/(K{r}+L{r}),0.5))"
        ws[f"Q{r}"] = f"=IF({vazio},NA(),1-P{r})"
        ws[f"R{r}"] = (f"=IF({vazio},NA(),MAX(0,E{r}-F{r})*IF(G{r}>0,H{r}/G{r},COMPAREC_PAD)"
                       f"*IF(I{r}>0,J{r}/I{r},VALIDOS_PAD))")
        ws[f"S{r}"] = f"=IF({vazio},NA(),IF(R{r}>0,0.5-(K{r}-L{r})/(2*R{r}),0))"
        ws[f"T{r}"] = formula_prob(r)
        ws[f"U{r}"] = f"=IF({vazio},NA(),1-T{r})"
        for ch in "NOPQSTU":
            ws[f"{ch}{r}"].number_format = "0.00%"
        ws[f"R{r}"].number_format = "#,##0"
        for ch in "CDEFGHIJKL":
            ws[f"{ch}{r}"].number_format = "#,##0"
    # esconde os #N/D das linhas ainda vazias (o #N/D faz o gráfico ignorar a linha)
    ws.conditional_formatting.add(f"N2:U{HIST_LINHAS + 1}",
                                  FormulaRule(formula=["ISNA(N2)"], font=Font(color="FFFFFF")))
    ws.freeze_panes = "C2"
    return ws


def kpi(ws, cel_rot, cel_val, rotulo, formula, fmt, cor=AZUL):
    ws.merge_cells(cel_rot)
    ws.merge_cells(cel_val)
    a = ws[cel_rot.split(":")[0]]
    a.value = rotulo
    a.font = Font(size=9, bold=True, color="595959")
    a.alignment = CENTRO
    a.fill = FUNDO_CINZA
    b = ws[cel_val.split(":")[0]]
    b.value = formula
    b.font = Font(size=16, bold=True, color=cor)
    b.alignment = CENTRO
    b.number_format = fmt
    b.fill = FUNDO_CINZA


def V(campo):
    return f'IFERROR(INDEX(Dados!$B$2:$B$40,MATCH("{campo}",Dados!$A$2:$A$40,0)),0)'


def aba_painel(wb, ws_hist):
    ws = wb.create_sheet("Painel", 0)
    ws.sheet_view.showGridLines = False
    largura(ws, {ch: 13 for ch in "ABCDEFGHIJKLMN"})
    faixa_titulo(ws, "A1:N1", "Apuração 2º turno — acompanhamento ao vivo (TSE)")
    ws.row_dimensions[1].height = 36
    ws.merge_cells("A2:N2")
    ws["A2"] = ('="Cargo: "&' + V("cargo") + '&"   ·   Abrangência: "&UPPER(' + V("cdabr") +
                ')&"   ·   Última totalização: "&' + V("dt") + '&" "&' + V("ht") +
                '&"   ·   Fonte: "&Config!B12')
    ws["A2"].font = Font(size=9, color="595959")

    andamento = (f'=IF({V("tf")}="s","Totalização final",IF({V("and")}="f","Finalizada",'
                 f'IF({V("and")}="p","Em andamento","Não iniciada")))')
    kpi(ws, "A4:B4", "A5:B6", "ANDAMENTO", andamento, "@")
    kpi(ws, "C4:D4", "C5:D6", "SEÇÕES TOTALIZADAS",
        f'=IFERROR({V("st")}/{V("ts")},0)', "0.00%")
    kpi(ws, "E4:F4", "E5:F6", "ELEITORADO APURADO",
        f'=IFERROR({V("est")}/{V("te")},0)', "0.00%")
    kpi(ws, "G4:H4", "G5:H6", "COMPARECIMENTO",
        f'=IFERROR({V("c")}/{V("esi")},0)', "0.00%")
    kpi(ws, "I4:J4", "I5:J6", "BRANCOS",
        f'=IFERROR({V("vb")}/{V("tv")},0)', "0.00%")
    kpi(ws, "K4:L4", "K5:L6", "NULOS",
        f'=IFERROR({V("tvn")}/{V("tv")},0)', "0.00%")
    kpi(ws, "M4:N4", "M5:N6", "MAT. DEFINIDO (TSE)",
        f'=IF({V("md")}="e","Eleito definido",IF({V("md")}="s","2º turno","Não"))', "@")

    # Tabela de candidatos
    cab = ["Nº", "Candidato", "", "Partido", "Votos", "", "% válidos", "Prob. vitória", "", "Situação TSE"]
    for j, t in enumerate(cab, start=1):
        c = ws.cell(8, j, t)
        c.font = CAB
        c.fill = FUNDO_AZUL
        c.alignment = CENTRO
    ws.merge_cells("B8:C8")
    ws.merge_cells("E8:F8")
    ws.merge_cells("H8:I8")
    ws.merge_cells("J8:K8")
    for i, (r, cor) in enumerate([(9, COR1), (10, COR2)]):
        dr = 2 + i
        ws.merge_cells(f"B{r}:C{r}")
        ws.merge_cells(f"E{r}:F{r}")
        ws.merge_cells(f"H{r}:I{r}")
        ws.merge_cells(f"J{r}:K{r}")
        ws[f"A{r}"] = f"=Dados!E{dr}"
        ws[f"B{r}"] = f"=Dados!F{dr}"
        ws[f"D{r}"] = f"=Dados!G{dr}"
        ws[f"E{r}"] = f"=Dados!H{dr}"
        ws[f"G{r}"] = "=IFERROR(E{0}/($E$9+$E$10),0)".format(r)
        ws[f"H{r}"] = "=P_CAND1" if i == 0 else "=1-P_CAND1"
        ws[f"J{r}"] = (f'=IF(Dados!J{dr}<>"",Dados!J{dr},IF(Dados!I{dr}="s","Eleito",'
                       f'IF(Dados!E{dr}="","","Em apuração")))')
        for col in "ABDEGHJ":
            c = ws[f"{col}{r}"]
            c.font = Font(bold=True, size=13, color=cor)
            c.alignment = CENTRO
            c.border = BORDA
        ws[f"E{r}"].number_format = "#,##0"
        ws[f"G{r}"].number_format = "0.00%"
        ws[f"H{r}"].number_format = "0.0%"
        ws.row_dimensions[r].height = 24
    ws.conditional_formatting.add("H9:I10", CellIsRule(operator="greaterThanOrEqual",
                                                       formula=["0.95"],
                                                       fill=PatternFill("solid", fgColor="E2EFDA")))

    # Cálculos da situação atual (mesmo modelo do Histórico)
    ws["L8"] = "Situação matemática"
    ws["L8"].font = CAB
    ws["L8"].fill = FUNDO_AZUL
    ws.merge_cells("L8:N8")
    calc = [
        (9, "Diferença (votos)", "=ABS(E9-E10)", "#,##0"),
        (10, "Válidos restantes (estim.)",
         f'=MAX(0,{V("te")}-{V("est")})*IF({V("esi")}>0,{V("c")}/{V("esi")},COMPAREC_PAD)'
         f'*IF({V("tv")}>0,{V("vv")}/{V("tv")},VALIDOS_PAD)', "#,##0"),
        (11, "2º colocado precisa dos restantes",
         '=IF(N10>0,0.5+N9/(2*N10),IF(N9>0,"Impossível",0.5))', "0.0%"),
        (12, "Leitura",
         f'=IF(OR({V("tf")}="s",{V("md")}="e"),"Definido pelo TSE",IF(N11>1,"Virada impossível",'
         'IF(N11>0.6,"Virada muito difícil",IF(N11>0.53,"Virada difícil","Disputa aberta"))))', "@"),
    ]
    for r, rot, f, fmt in calc:
        ws.merge_cells(f"L{r}:M{r}")
        ws[f"L{r}"] = rot
        ws[f"L{r}"].font = Font(size=9, bold=True, color="595959")
        ws[f"L{r}"].alignment = Alignment(wrap_text=True, vertical="center")
        ws[f"N{r}"] = f
        ws[f"N{r}"].number_format = fmt
        ws[f"N{r}"].font = NEGRITO
        ws[f"N{r}"].alignment = CENTRO
        for col in "LMN":
            ws[f"{col}{r}"].border = BORDA

    # probabilidade atual do candidato 1 (célula auxiliar nomeada)
    ws["P9"] = "aux: prob. cand.1"
    ws["Q9"] = (f'=IF(OR({V("tf")}="s",{V("md")}="e"),IF(Dados!I2="s",1,IF(Dados!I3="s",0,'
                'IF(E9>E10,1,0))),IF(N10<=0,IF(E9>E10,1,IF(E9<E10,0,0.5)),'
                '1-NORMDIST(0.5-(E9-E10)/(2*N10),IF(E9+E10>0,E9/(E9+E10),0.5)+AJUSTE,SIGMA,TRUE)))')
    ws["P9"].font = Font(color="BFBFBF", size=8)
    ws["Q9"].font = Font(color="BFBFBF", size=8)
    wb.defined_names["P_CAND1"] = DefinedName("P_CAND1", attr_text="Painel!$Q$9")

    ws["A13"] = ('Probabilidade = chance de cada candidato terminar à frente, supondo que a fatia dele nas '
                 'urnas ainda não apuradas siga a fatia atual ± a incerteza definida em Config (distribuição '
                 'normal). É estimativa, não resultado oficial.')
    ws.merge_cells("A13:N13")
    ws["A13"].font = Font(italic=True, size=9, color="7F7F7F")
    ws["A13"].alignment = QUEBRA
    ws.row_dimensions[13].height = 28

    ultima = HIST_LINHAS + 1
    xs = Reference(ws_hist, min_col=14, min_row=2, max_row=ultima)  # % seções

    def dispersao(titulo, cols, eixo_y, ancora, y_max=1.0, y_min=0.0):
        ch = ScatterChart()
        ch.title = titulo
        ch.style = 13
        ch.height = 9
        ch.width = 17
        ch.x_axis.title = "% de seções totalizadas"
        ch.y_axis.title = eixo_y
        ch.x_axis.number_format = "0%"
        ch.y_axis.number_format = "0%"
        ch.x_axis.scaling.min = 0
        ch.x_axis.scaling.max = 1
        ch.y_axis.scaling.min = y_min
        ch.y_axis.scaling.max = y_max
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        ch.legend.position = "b"
        for col, cor, nome in cols:
            ys = Reference(ws_hist, min_col=col, min_row=2, max_row=ultima)
            s = Series(ys, xs, title_from_data=False, title=nome)
            s.graphicalProperties.line.solidFill = cor
            s.graphicalProperties.line.width = 28000
            s.marker.symbol = "circle"
            s.marker.size = 4
            s.marker.graphicalProperties.solidFill = cor
            s.marker.graphicalProperties.line.solidFill = cor
            s.smooth = False
            ch.series.append(s)
        ws.add_chart(ch, ancora)

    dispersao("Probabilidade de vitória × apuração",
              [(20, COR1, "Lula (13)"), (21, COR2, "Flávio Bolsonaro (22)")],
              "Probabilidade", "A15")
    dispersao("% dos votos válidos × apuração",
              [(16, COR1, "Lula (13)"), (17, COR2, "Flávio Bolsonaro (22)")],
              "% válidos", "H15", y_max=0.65, y_min=0.35)

    bar = BarChart()
    bar.type = "bar"
    bar.title = "Votos totalizados"
    bar.style = 10
    bar.height = 7
    bar.width = 17
    bar.add_data(Reference(ws, min_col=5, min_row=9, max_row=10), titles_from_data=False)
    bar.set_categories(Reference(ws, min_col=2, min_row=9, max_row=10))
    bar.legend = None
    bar.y_axis.number_format = "#,##0"
    bar.y_axis.delete = False
    bar.x_axis.delete = False
    bar.y_axis.scaling.min = 0
    for idx, cor in ((0, COR1), (1, COR2)):
        pt = DataPoint(idx=idx)
        pt.graphicalProperties.solidFill = cor
        bar.series[0].dPt.append(pt)
    bar.dataLabels = DataLabelList()
    bar.dataLabels.showVal = True
    bar.dataLabels.showSerName = False
    bar.dataLabels.showCatName = False
    bar.dataLabels.showLegendKey = False
    ws.add_chart(bar, "A34")

    bar2 = BarChart()
    bar2.type = "bar"
    bar2.grouping = "percentStacked"
    bar2.overlap = 100
    bar2.title = "Seções: totalizadas × faltando"
    bar2.style = 10
    bar2.height = 7
    bar2.width = 17
    ws["P11"] = "Totalizadas"
    ws["Q11"] = f"={V('st')}"
    ws["P12"] = "Faltando"
    ws["Q12"] = f"=MAX(0,{V('ts')}-{V('st')})"
    for c in ("P11", "Q11", "P12", "Q12"):
        ws[c].font = Font(color="BFBFBF", size=8)
    for row, cor in ((11, AZUL), (12, "BFBFBF")):
        bar2.add_data(Reference(ws, min_col=17, min_row=row, max_row=row), titles_from_data=False)
        bar2.series[-1].graphicalProperties.solidFill = cor
    bar2.series[0].tx = SeriesLabel(v="Totalizadas")
    bar2.series[1].tx = SeriesLabel(v="Faltando")
    bar2.legend.position = "b"
    bar2.y_axis.number_format = "0%"
    bar2.y_axis.delete = False
    bar2.x_axis.delete = True
    ws.add_chart(bar2, "H34")
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = "A1:N48"
    return ws


def aba_regras(wb):
    ws = wb.create_sheet("Regras TSE")
    largura(ws, {"A": 16, "B": 110})
    faixa_titulo(ws, "A1:B1", "Regras do arquivo EA20 (resultado unificado) usadas na planilha")
    ws.row_dimensions[1].height = 32
    regras = [
        ("Arquivo", "EA20 — resultado unificado: parte fixa (candidatos, partidos, seções, eleitorado) + "
                    "parte variável (votos), atualizada a cada totalização."),
        ("Nome", "br-c<cargo>-e<eleição>-u.json (nacional) · <uf>-c<cargo>-e<eleição>-u.json (estadual) · "
                 "<uf><município 5 díg.>-c…-u.json (municipal) · <uf><mun>-z<zona 4 díg.>-c…-u.json (zona)."),
        ("Cargos", "0001 Presidente (BR) · 0003 Governador (UF) · 0005 Senador · 0006 Dep. Federal · "
                   "0007 Dep. Estadual · 0008 Dep. Distrital · 0011 Prefeito · 0013 Vereador. "
                   "2º turno: apenas Presidente e Governador."),
        ("t", "Turno: 1 = primeiro, 2 = segundo."),
        ("f", "Fase: s = simulado, o = oficial. Confira se está 'o' no dia da eleição."),
        ("dt / ht", "Data e hora da totalização — a planilha grava um ponto no histórico a cada mudança."),
        ("and", "Andamento: n = não iniciada, p = em andamento, f = finalizada. "
                "'Finalizada' ≠ 'totalização final' (tf)."),
        ("tf", "Totalização final: s/n. É quando o juiz eleitoral finaliza a eleição."),
        ("md", "Matematicamente definido (só Presidente/Governador/Prefeito, enquanto tf = n): "
               "e = eleito, s = segundo turno, n = não definido. Quando 'e', a planilha fixa 100%/0%."),
        ("dv", "Divulga votação: se 'n', os votos vêm zerados (regra de divulgação do TSE — "
               "a divulgação começa às 17h de Brasília)."),
        ("cand.e / st", "e = s se eleito (ou foi ao 2º turno). st (Eleito, Não eleito, 2º turno…) só é "
                        "preenchido após a totalização final."),
        ("vap / pvap", "Votos computados do candidato e % sobre os votos a votáveis concorrentes. "
                       "A planilha recalcula % sobre a soma dos dois candidatos (= válidos no 2º turno)."),
        ("Seções (s)", "ts = st (totalizadas) + snt (não totalizadas); st = si (instaladas) + sni; "
                       "si = sa (apuradas) + sna."),
        ("Eleitores (e)", "te = est + esnt; est = esi + esni; esi = c (comparecimento) + a (abstenção) "
                          "= esa + esna."),
        ("Votos (v)", "tv = vb + vn + vnt + van + vansj + vv; tvn = vn + vnt (nulos + nulos técnicos); "
                      "vv = vnom + vl."),
        ("Percentuais", "Cada % vem em 2 campos: 2 casas (texto '12,34') e 9 casas (pstn, pvapn…). "
                        "Arredondamento: ≥5 sobe, <5 mantém. A planilha calcula % pelas quantidades."),
        ("Ordem", "seq = ordem do candidato (muda conforme votos). A planilha ordena pelo número de urna, "
                  "para o cand. 1 e o cand. 2 não trocarem de lugar no histórico."),
    ]
    ws["A3"], ws["B3"] = "Campo", "Regra / uso"
    for c in (ws["A3"], ws["B3"]):
        c.font = CAB
        c.fill = FUNDO_AZUL
    for i, (a, b) in enumerate(regras, start=4):
        ws.cell(i, 1, a).font = NEGRITO
        ws.cell(i, 2, b).alignment = QUEBRA
        ws.cell(i, 1).alignment = Alignment(vertical="top")
        ws.row_dimensions[i].height = 32


def aba_como_usar(wb):
    ws = wb.create_sheet("Como usar")
    largura(ws, {"A": 4, "B": 120})
    faixa_titulo(ws, "A1:B1", "Como conectar ao TSE e acompanhar ao vivo")
    ws.row_dimensions[1].height = 32
    passos = [
        ("1. Antes de tudo", None),
        ("", "Em Config, confira a URL (B12). O padrão aponta para o 2º turno 2026 (eleição 6258, Presidente). "
             "Para governador, troque Abrangência para a UF (ex.: sp) e Cargo para 3. Se o TSE usar outro "
             "caminho, cole a URL exata em Config!B8 (veja o link do JSON na aba Documentos de "
             "resultados.tse.jus.br)."),
        ("", "Caminho mais simples: o painel web da pasta painel-ao-vivo (dois cliques, já conectado). "
             "Esta planilha é a alternativa para quem quer os números no Excel ou no Google Planilhas."),
        ("2A. Excel (Microsoft 365 / 2016+) — Power Query", None),
        ("", "Arquivo > Opções > Central de Confiabilidade/Privacidade > Power Query: marque "
             "'Ignorar os níveis de privacidade' (necessário porque a consulta usa a URL da célula Config!B12)."),
        ("", "Apague (tecla Delete) Dados!A1:B40, Dados!E1:K20 e Histórico!A1:L1: as consultas recriam esses dados e cabeçalhos."),
        ("", "Dados > Obter Dados > De Outras Fontes > Consulta Nula. Página Inicial > Editor Avançado. "
             "Cole o bloco 'TSE' do arquivo powerquery_tse.m e renomeie a consulta para TSE."),
        ("", "Repita para Resumo, Candidatos e tbHistorico (os nomes precisam ser exatamente esses)."),
        ("", "Fechar e Carregar Para…: TSE = Somente Conexão; Resumo = Tabela em Dados!A1; "
             "Candidatos = Tabela em Dados!E1; tbHistorico = Tabela em Histórico!A1."),
        ("", "Dados > Consultas e Conexões > botão direito em cada consulta > Propriedades: marque "
             "'Atualizar a cada 1 minuto' e 'Atualizar dados ao abrir o arquivo'. Pronto: o Painel e os "
             "gráficos atualizam sozinhos."),
        ("", "Salve como .xlsx normal. Mantenha o arquivo aberto durante a apuração."),
        ("2B. Google Planilhas — Apps Script", None),
        ("", "Envie este .xlsx ao Google Drive e abra com Planilhas Google (Arquivo > Salvar como Planilha Google)."),
        ("", "Extensões > Apps Script. Cole o conteúdo de tse_apps_script.gs, salve, recarregue a planilha."),
        ("", "Menu 'TSE ao vivo' > 'Atualizar agora' (autorize na 1ª vez) > 'Iniciar ao vivo (1 min)'."),
        ("", "Para parar: 'Parar ao vivo'. O histórico cresce uma linha a cada nova totalização do TSE."),
        ("3. Lendo o painel", None),
        ("", "Prob. vitória: chance estimada de terminar à frente considerando os votos que faltam. "
             "Ajuste a incerteza em Config!B16 (maior = mais cautela com viradas)."),
        ("", "'2º colocado precisa dos restantes': % dos votos válidos ainda não apurados que o candidato "
             "atrás precisa obter para virar. Acima de 100% = virada impossível."),
        ("", "Quando o TSE marcar md = 'e' (matematicamente definido) ou tf = 's', a planilha fixa 100%/0%."),
        ("", "Teste antes do dia 25/10 usando o ambiente simulado do TSE (endereço em Config!C3) colando a "
             "URL completa do simulado em Config!B8."),
    ]
    r = 3
    for titulo, texto in passos:
        if texto is None:
            ws.cell(r, 1, titulo).font = Font(bold=True, size=13, color=AZUL)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        else:
            ws.cell(r, 1, "•")
            ws.cell(r, 2, texto).alignment = QUEBRA
            ws.row_dimensions[r].height = 18 * max(1, math.ceil(len(texto) / 125))
        r += 1

    r += 1
    ws.cell(r, 1, "Código Power Query (mesmo conteúdo de powerquery_tse.m)").font = Font(bold=True, size=13, color=AZUL)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    r += 1
    for linha in (AQUI / "powerquery_tse.m").read_text(encoding="utf-8").splitlines():
        c = ws.cell(r, 2, linha)
        c.font = Font(name="Consolas", size=9)
        r += 1


def main():
    saida = Path(sys.argv[1]) if len(sys.argv) > 1 else AQUI / "apuracao_2turno.xlsx"
    wb = Workbook()
    aba_config(wb)
    ws_dados = aba_dados(wb)
    ws_hist = aba_historico(wb)
    preencher_inicial(ws_dados)
    aba_painel(wb, ws_hist)
    aba_regras(wb)
    aba_como_usar(wb)
    wb.active = 0
    wb.save(saida)
    print(f"Gerado: {saida}")


if __name__ == "__main__":
    main()
